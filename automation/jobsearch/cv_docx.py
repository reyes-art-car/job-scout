"""Generación del CV adaptado en .docx."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import HDR_TEAL, Profile
from .offer import JobOffer

EXPERIENCE_HEADER = "FREELANCE  |  Desarrolladora Java/Angular"
EXPERIENCE_DATES = "Marzo 2026 - Actualidad"
PROJECT_LINE = "CRM de gestion comercial · Desarrollo, despliegue y evolucion continua"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _spacing(paragraph, before: int = 0, after: int = 0) -> None:
    element = OxmlElement("w:spacing")
    element.set(qn("w:before"), str(before))
    element.set(qn("w:after"), str(after))
    paragraph._p.get_or_add_pPr().append(element)


def _section_title(doc, title: str) -> None:
    paragraph = doc.add_paragraph()
    _spacing(paragraph, before=80, after=20)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(HDR_TEAL)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attribute, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", HDR_TEAL)):
        bottom.set(qn(f"w:{attribute}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _bullet(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    _spacing(paragraph, before=20, after=20)
    paragraph.add_run(text).font.size = Pt(8.5)


def _line(doc, parts: list[tuple[str, bool]], before: int = 30, after: int = 10) -> None:
    paragraph = doc.add_paragraph()
    _spacing(paragraph, before=before, after=after)
    for text, bold in parts:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(8.5)


class CvDocumentBuilder:
    def __init__(self, profile: Profile, output_dir: Path) -> None:
        self._profile = profile
        self._output_dir = output_dir.resolve()

    def output_path(self, filename: str) -> Path:
        """Ruta de salida dentro del directorio de CVs; se rechaza cualquier nombre que intente salir de él."""
        target = (self._output_dir / filename).resolve()
        if target.parent != self._output_dir:
            raise ValueError(f"Nombre de fichero no permitido: {filename!r}")
        return target

    def build(self, offer: JobOffer) -> Path:
        if offer.cv_file is None:
            raise ValueError("La oferta no tiene nombre de CV")
        target = self.output_path(offer.cv_file)
        self._output_dir.mkdir(parents=True, exist_ok=True)
        self._compose(offer).save(target)
        return target

    def _compose(self, offer: JobOffer):
        profile = self._profile
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Cm(1.2)
            section.left_margin = section.right_margin = Cm(1.8)

        self._header(doc, offer)
        doc.add_paragraph()

        _section_title(doc, "PERFIL PROFESIONAL")
        summary = doc.add_paragraph()
        _spacing(summary, before=30, after=30)
        summary.add_run(offer.cv_summary).font.size = Pt(8.5)

        _section_title(doc, "EXPERIENCIA Y PROYECTOS")
        employer = doc.add_paragraph()
        _spacing(employer, before=30, after=0)
        title_run = employer.add_run(EXPERIENCE_HEADER)
        title_run.bold = True
        title_run.font.size = Pt(9)
        employer.add_run("                                                       " + EXPERIENCE_DATES).font.size = Pt(
            8.5
        )
        project = doc.add_paragraph()
        _spacing(project, before=0, after=10)
        project_run = project.add_run(PROJECT_LINE)
        project_run.italic = True
        project_run.font.size = Pt(8.5)
        project_run.font.color.rgb = _rgb(HDR_TEAL)
        order = offer.cv_bullets or list(range(len(profile.bullets)))
        for index in order:
            _bullet(doc, profile.bullets[index])

        _section_title(doc, "COMPETENCIAS TECNICAS")
        groups = profile.skills.get(offer.cv_skills_order) or profile.skills["backend_first"]
        for group in groups:
            _line(doc, [(f"{group.label}: ", True), (group.content, False)], before=20)

        _section_title(doc, "FORMACION")
        _line(
            doc,
            [
                ("Tecnica Superior en Desarrollo de Aplicaciones Web (DAW)", True),
                ("  |  MEDAC Sevilla  |  2024-2026", False),
            ],
        )
        _section_title(doc, "IDIOMAS")
        _line(doc, [("Ingles C1", True), (" - acreditado por la Universidad de Cambridge", False)])
        return doc

    def _header(self, doc, offer: JobOffer) -> None:
        table = doc.add_table(rows=0, cols=1)
        table.style = "Table Grid"
        for text, size, bold, before, after in (
            (self._profile.name, 18, True, 60, 20),
            (offer.cv_title, 10, False, 0, 20),
            (self._profile.contact_line, 8.5, False, 0, 60),
        ):
            cell = table.add_row().cells[0]
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:color"), "auto")
            shade.set(qn("w:fill"), HDR_TEAL)
            cell._tc.get_or_add_tcPr().append(shade)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = cell.paragraphs[0].add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _spacing(cell.paragraphs[0], before=before, after=after)
