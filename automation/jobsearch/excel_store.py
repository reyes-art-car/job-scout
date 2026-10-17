"""Excel de ofertas (lectura de URLs conocidas y alta de filas nuevas)."""

from collections.abc import Iterable
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import EXCEL_COLUMNS, XL_LIGHT, XL_TEAL
from .offer import JobOffer

_COLUMN_WIDTHS = [8, 10, 12, 30, 20, 12, 45, 14, 20, 22, 18, 20, 18, 45, 35, 50, 45, 60, 20, 40, 40]
_URL_COLUMN = 7
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
GREEN, YELLOW, RED = "C6EFCE", "FFEB9C", "FFC7CE"


def _score_color(compatibility: int) -> str:
    return GREEN if compatibility >= 85 else YELLOW if compatibility >= 70 else RED


def _set_text(cell, value: str) -> None:
    """Escribe siempre como TEXTO: un valor que empiece por «=» no debe convertirse en fórmula de Excel
    (inyección de fórmulas a partir de contenido de un anuncio)."""
    cell.value = value
    if value:
        cell.data_type = "s"


class OfferWorkbook:
    def __init__(self, path: Path) -> None:
        self._path = path

    def known_urls_and_last_id(self) -> tuple[set[str], int]:
        if not self._path.exists():
            return set(), 0
        workbook = openpyxl.load_workbook(self._path, read_only=True)
        urls: set[str] = set()
        last = 0
        try:
            sheet = workbook.active
            if sheet is None:
                return urls, last
            for row in sheet.iter_rows(min_row=2, values_only=True):
                code, url = row[0], row[_URL_COLUMN - 1] if len(row) >= _URL_COLUMN else None
                if code:
                    digits = str(code).replace("OF", "")
                    if digits.isdigit():
                        last = max(last, int(digits))
                if url:
                    urls.add(str(url).strip())
        finally:
            workbook.close()
        return urls, last

    def append(self, offers: Iterable[JobOffer]) -> int:
        workbook = openpyxl.load_workbook(self._path) if self._path.exists() else self._new_workbook()
        sheet = workbook.active
        if sheet is None:
            raise RuntimeError("El libro de Excel no tiene hoja activa")
        added = 0
        for offer in offers:
            self._write_row(sheet, sheet.max_row + 1, offer)
            added += 1
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_COLUMNS))}{sheet.max_row}"
        workbook.save(self._path)
        return added

    @staticmethod
    def _new_workbook() -> openpyxl.Workbook:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Ofertas"
        for column, name in enumerate(EXCEL_COLUMNS, 1):
            cell = sheet.cell(row=1, column=column, value=name)
            cell.fill = PatternFill("solid", fgColor=XL_TEAL)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = _BORDER
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "A2"
        for column, width in enumerate(_COLUMN_WIDTHS, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        return workbook

    @staticmethod
    def _write_row(sheet: Worksheet, row: int, offer: JobOffer) -> None:
        values = [
            offer.id, offer.priority, f"{offer.compatibility}%", offer.title, offer.company, offer.portal,
            offer.url, offer.date, offer.location, offer.modality, offer.contract, offer.salary,
            offer.experience, offer.mandatory_tech, offer.desirable_tech, offer.meets, offer.missing,
            offer.score_reason, offer.status, offer.cv_file or "", offer.notes,
        ]  # fmt: skip
        background = XL_LIGHT if row % 2 == 0 else "FFFFFF"
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column)
            _set_text(cell, value)
            cell.fill = PatternFill("solid", fgColor=background)
            cell.font = Font(size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = _BORDER
        link = sheet.cell(row=row, column=_URL_COLUMN)
        link.hyperlink = offer.url  # la URL ya está validada como http/https
        link.font = Font(size=9, color="0563C1", underline="single")
        score = sheet.cell(row=row, column=3)
        score.fill = PatternFill("solid", fgColor=_score_color(offer.compatibility))
        score.font = Font(bold=True, size=10)
        score.alignment = Alignment(horizontal="center", vertical="top")
